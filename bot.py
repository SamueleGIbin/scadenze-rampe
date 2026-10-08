import os
import time
import datetime as dt
from zoneinfo import ZoneInfo
import requests
import openpyxl

TZ = ZoneInfo("Europe/Rome")
MAX = 345 * 60
COLORS = {1: "Good", 2: "Warning", 3: "Attention"}

def secs(v):
    if isinstance(v, (dt.datetime, dt.time)):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, dt.timedelta):
        return int(v.total_seconds()) % 86400
    return None

def nxt(s, now):
    d = now.replace(hour=0, minute=0, second=0, microsecond=0) + dt.timedelta(seconds=s)
    return d if d >= now else d + dt.timedelta(days=1)

def load(now):
    ws = openpyxl.load_workbook("data.xlsx", data_only=True).active
    ref = secs(ws["G2"].value)
    tasks = []
    for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        ramp, dest, dep, frt, closed = r[0], r[1], secs(r[2]), secs(r[3]), r[4]
        if not (isinstance(ramp, (int, float)) and dep is not None and not closed):
            continue
        if frt is not None and ref is not None:
            s, verb = (ref + frt) % 86400, "girare"
        else:
            s, verb = dep, "chiudere"
        d = nxt(s, now)
        tasks.append({"deadline": d, "text": f"Baia {int(ramp)} - {dest} - da {verb} {d:%H:%M}", "sent": 0})
    tasks.sort(key=lambda t: t["deadline"])
    return tasks

def step(tasks, now):
    lines = []
    for t in tasks:
        m = (t["deadline"] - now).total_seconds() / 60
        lvl = 3 if m <= 10 else 2 if m <= 15 else 1 if m <= 30 else 0
        if m > 0 and lvl > t["sent"]:
            t["sent"] = lvl
            lines.append((COLORS[lvl], t["text"]))
    return lines

def send(lines):
    body = [{"type": "TextBlock", "text": x, "color": c, "weight": "Bolder", "wrap": True} for c, x in lines]
    requests.post(
        os.environ["TEAMS_URL"],
        json={
            "type": "message",
            "attachments": [{
                "contentType": "application/vnd.microsoft.card.adaptive",
                "content": {
                    "$schema": "http://adaptivecards.io/schemas/adaptive-card.json",
                    "type": "AdaptiveCard",
                    "version": "1.4",
                    "body": body,
                },
            }],
        },
        timeout=30,
    ).raise_for_status()

def run():
    t0 = dt.datetime.now(TZ)
    tasks = [t for t in load(t0) if (t["deadline"] - t0).total_seconds() <= MAX]
    while tasks and (dt.datetime.now(TZ) - t0).total_seconds() < MAX:
        now = dt.datetime.now(TZ)
        lines = step(tasks, now)
        if lines:
            send(lines)
        tasks = [t for t in tasks if t["deadline"] > now]
        time.sleep(30)

if __name__ == "__main__":
    run()
